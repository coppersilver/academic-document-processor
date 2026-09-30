from pathlib import Path
from pptx import Presentation

def extract_pptx_text(file_path: str | Path) -> str:
    """Extracts text and speaker notes from PowerPoint (.pptx) presentations."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"PPTX not found: {file_path}")

    prs = Presentation(str(path))
    slides_text: list[str] = []

    for i, slide in enumerate(prs.slides):
        slide_parts: list[str] = []

        # Extract text from shapes
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_parts.append(text)
            elif shape.has_table:
                for row in shape.table.rows:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(row_cells):
                        slide_parts.append(" | ".join(row_cells))

        # Check for notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_parts.append(f"[Speaker Notes: {notes}]")

        if slide_parts:
            slides_text.append(f"--- Slide {i + 1} ---\n" + "\n".join(slide_parts))

    return "\n\n".join(slides_text)
