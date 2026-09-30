from pathlib import Path
import docx

def extract_docx_text(file_path: str | Path) -> str:
    """Extracts text and tables from Word (.docx) documents."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX not found: {file_path}")

    doc = docx.Document(str(path))
    content_parts: list[str] = []

    # Extract paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            content_parts.append(text)

    # Extract tables
    for t_idx, table in enumerate(doc.tables):
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(row_cells):
                table_rows.append(" | ".join(row_cells))
        if table_rows:
            content_parts.append(f"\n[Table {t_idx + 1}]\n" + "\n".join(table_rows))

    return "\n\n".join(content_parts)
